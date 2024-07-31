import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from docx import Document

def create_html_from_docx(docx_file):
    document = Document(docx_file)
    html_content = ""

    for paragraph in document.paragraphs:
        if paragraph.runs:
            html_paragraph = ""
            for run in paragraph.runs:
                text = run.text
                if run.bold:
                    text = f"<b>{text}</b>"
                if run.italic:
                    text = f"<i>{text}</i>"
                if run.underline:
                    text = f"<u>{text}</u>"
                html_paragraph += text
            html_content += f"<p>{html_paragraph}</p>"
        else:
            html_content += "<p></p>"

    return html_content

def create_email_body_from_file(
    employee_name,
    employee_id,
    registration_date_time,
    supervisor_name,
    similarity_score,
    docx_file="email_body.docx"
):
    try:
        html_body = create_html_from_docx(docx_file)
        
        # Replace placeholders with actual values
        html_body = html_body.replace("[Employee Name]", employee_name)
        html_body = html_body.replace("[Employee ID]", employee_id)
        html_body = html_body.replace("[Date and Time]", registration_date_time)
        html_body = html_body.replace("[Supervisor Name]", supervisor_name)
        html_body = html_body.replace("[Score Indicating Similarity Level]", similarity_score)

        return html_body
    except Exception as e:
        print(f"Failed to create email body: {str(e)}")
        return None

def send_email(to_address, subject, cc, body):
    """
    Sends an email.

    Args:
        to_address (list): List of email addresses of the recipients.
        subject (str): Subject of the email.
        cc (str): Comma-separated list of CC recipients.
        body (str): Body content of the email.
    """
    try:
        smtp_host = ""  # replace with actual SMTP host
        smtp_port = 587  # replace with actual SMTP port
        smtp_user = ""  # replace with actual SMTP user
        smtp_password = ""  # replace with actual SMTP password



        s = smtplib.SMTP(smtp_host, smtp_port)
        s.starttls()
        s.login(smtp_user, smtp_password)

        message = MIMEMultipart("alternative")
        message["From"] = smtp_user
        message["To"] = ", ".join(to_address)  # join list of addresses into a single string
        message["CC"] = cc
        message["Subject"] = subject
        message.attach(MIMEText(body, "html"))

        # Combine to and cc recipients
        recipients = to_address + cc.split(",") if cc else to_address

        s.sendmail(smtp_user, recipients, message.as_string())
        s.quit()
        print("Email sent successfully!")
        return True
    except Exception as e:
        print(f"Failed to send email: {str(e)}")
        return False

# Example usage:

employee_name = "John Doe"
employee_id = "12345"
registration_date_time = "2024-07-30 10:00 AM"
supervisor_name = "Jane Smith"
similarity_score = "0.95"

email_body = create_email_body_from_file(
    employee_name,
    employee_id,
    registration_date_time,
    supervisor_name,
    similarity_score,
)

if email_body:
    to_address = ["shreyanshdewangan4299@gmail.com", "shreyansh.d@adcuratio.com"]
    send_email(
        to_address=to_address,
        subject="Employee Registration",
        cc="cc@example.com",
        body=email_body,
    )
